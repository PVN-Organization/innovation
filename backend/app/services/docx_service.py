"""Non-blocking DOCX export using docxtpl + python-docx.

Simple variables are rendered via docxtpl.  The author table uses
``__AUTHOR_*__`` placeholders that are replaced by python-docx after
rendering, which avoids the fragile ``{%tr%}`` tag processing in docxtpl.
"""

import asyncio
from copy import deepcopy
from datetime import UTC, datetime
from io import BytesIO
import json
from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate
from lxml import etree

TEMPLATE_PATH = (
    Path(__file__).resolve().parent.parent / "templates" / "mau_sang_kien.docx"
)

_PLACEHOLDER_MAP = {
    "__AUTHOR_ROLE__": "vaiTro",
    "__AUTHOR_NAME__": "hoTen",
    "__AUTHOR_TITLE__": "chucVu",
    "__AUTHOR_UNIT__": "donVi",
    "__AUTHOR_EMAIL__": "email",
}


def _parse_authors(data: dict[str, object]) -> list[dict[str, str]]:
    """Try structured JSON first; fall back to legacy flat fields."""
    raw = str(data.get("danh_sach_tac_gia", "") or "")
    if raw:
        try:
            parsed = json.loads(raw)
            if isinstance(parsed, list) and len(parsed) > 0:
                return parsed
        except (json.JSONDecodeError, TypeError):
            pass

    tac_gia = str(data.get("tac_gia", ""))
    if not tac_gia:
        return [{"vaiTro": "", "hoTen": "", "chucVu": "", "donVi": "", "email": ""}]

    authors: list[dict[str, str]] = [
        {
            "vaiTro": "Tác giả chính",
            "hoTen": tac_gia,
            "chucVu": "",
            "donVi": "",
            "email": "",
        },
    ]
    dong = str(data.get("dong_tac_gia", ""))
    for name in dong.split(";"):
        name = name.strip()
        if name:
            authors.append({
                "vaiTro": "Đồng tác giả",
                "hoTen": name,
                "chucVu": "",
                "donVi": "",
                "email": "",
            })
    return authors


def _populate_author_table(doc: Document, authors: list[dict[str, str]]) -> None:
    """Replace __AUTHOR_*__ placeholder row with one row per author."""
    table = doc.tables[0]
    rows = table.rows

    template_row_el = rows[1]._tr
    parent = template_row_el.getparent()

    for author in authors:
        new_row_el = deepcopy(template_row_el)
        for cell_el in new_row_el.findall(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc"
        ):
            for t_el in cell_el.iter(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
            ):
                text = t_el.text or ""
                for placeholder, key in _PLACEHOLDER_MAP.items():
                    if placeholder in text:
                        t_el.text = text.replace(placeholder, author.get(key, ""))
                        break
        parent.insert(list(parent).index(template_row_el), new_row_el)

    parent.remove(template_row_el)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _fix_newlines(doc: Document) -> None:
    """Convert literal ``\\n`` inside <w:t> elements into <w:br/> line breaks."""
    for t_el in doc.element.body.iter(f"{{{W_NS}}}t"):
        text = t_el.text or ""
        if "\n" not in text:
            continue

        parts = text.split("\n")
        run_el = t_el.getparent()

        t_el.text = parts[0]

        anchor = t_el
        for part in parts[1:]:
            br = etree.SubElement(run_el, f"{{{W_NS}}}br")
            run_el.insert(list(run_el).index(anchor) + 1, br)
            new_t = etree.SubElement(run_el, f"{{{W_NS}}}t")
            new_t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            new_t.text = part
            run_el.insert(list(run_el).index(br) + 1, new_t)
            anchor = new_t


def _render(context: dict[str, object], authors: list[dict[str, str]]) -> bytes:
    """Render template: docxtpl for text vars, python-docx for author table."""
    tpl = DocxTemplate(str(TEMPLATE_PATH))
    tpl.render(context)

    buf = BytesIO()
    tpl.save(buf)
    buf.seek(0)

    doc = Document(buf)
    _populate_author_table(doc, authors)
    _fix_newlines(doc)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _build_initiative_document(data: dict[str, object]) -> bytes:
    """Render from DB data (snake_case keys)."""
    authors = _parse_authors(data)
    now = datetime.now(tz=UTC)

    context = {
        "ten": data.get("ten", ""),
        "linhVuc": data.get("linh_vuc", ""),
        "thoiGian": data.get("thoi_gian", ""),
        "lyDo": data.get("ly_do", ""),
        "mucTieu": data.get("muc_tieu", ""),
        "thucTrang": data.get("thuc_trang", ""),
        "giaiPhap": data.get("giai_phap", ""),
        "cachThuc": data.get("cach_thuc", ""),
        "hieuQua": data.get("hieu_qua", ""),
        "tinhMoi": data.get("tinh_moi", ""),
        "nhanRong": data.get("nhan_rong", ""),
        "dd": now.strftime("%d"),
        "mm": now.strftime("%m"),
        "yyyy": now.strftime("%Y"),
    }
    return _render(context, authors)


def _build_preview_document(form_data: dict[str, object]) -> bytes:
    """Render from frontend form data (camelCase keys)."""
    raw_authors = str(form_data.get("danhSachTacGia", "") or "")
    authors: list[dict[str, str]] = []
    if raw_authors:
        try:
            parsed = json.loads(raw_authors)
            if isinstance(parsed, list):
                authors = parsed
        except (json.JSONDecodeError, TypeError):
            pass
    if not authors:
        authors = [{"vaiTro": "", "hoTen": "", "chucVu": "", "donVi": "", "email": ""}]

    now = datetime.now(tz=UTC)

    context = {
        "ten": form_data.get("ten", ""),
        "linhVuc": form_data.get("linhVuc", ""),
        "thoiGian": form_data.get("thoiGian", ""),
        "lyDo": form_data.get("lyDo", ""),
        "mucTieu": form_data.get("mucTieu", ""),
        "thucTrang": form_data.get("thucTrang", ""),
        "giaiPhap": form_data.get("giaiPhap", ""),
        "cachThuc": form_data.get("cachThuc", ""),
        "hieuQua": form_data.get("hieuQua", ""),
        "tinhMoi": form_data.get("tinhMoi", ""),
        "nhanRong": form_data.get("nhanRong", ""),
        "dd": now.strftime("%d"),
        "mm": now.strftime("%m"),
        "yyyy": now.strftime("%Y"),
    }
    return _render(context, authors)


async def generate_initiative_docx(data: dict[str, object]) -> bytes:
    """Generate DOCX from DB data without blocking."""
    return await asyncio.to_thread(_build_initiative_document, data)


async def generate_preview_docx(form_data: dict[str, object]) -> bytes:
    """Generate DOCX from form data without blocking."""
    return await asyncio.to_thread(_build_preview_document, form_data)
